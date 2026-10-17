import os
import uuid
from datetime import datetime
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, Body
from pydantic import BaseModel
try:
    from sqlalchemy.orm import Session
except ImportError:
    Session = object

from app.config import UPLOAD_DIR
from app.database import get_db
from app.models import SiteBlock, User
from app.auth import require_cadre

router = APIRouter(prefix="/api/slides", tags=["Canva 簡報與視覺編輯器 (Canva-like PPT Studio)"])

class SlidePage(BaseModel):
    id: Optional[str] = None
    template: str = "evaluation"
    theme: str = "amber"
    title: str = "四大宗旨活動成果評鑑"
    subtitle: str = "國際佛光會新竹青年分團"
    bullets: List[str] = []
    metrics: List[Dict[str, str]] = []
    badge: str = "三好四給・生活實踐"
    footer: str = "做好事・說好話・存好心"
    bg_image: Optional[str] = ""

class DeckExportRequest(BaseModel):
    deck_title: str = "新竹佛光青年分團簡報"
    slides: List[SlidePage]

@router.get("/get-deck")
def get_saved_deck(db: Session = Depends(get_db)):
    block = db.query(SiteBlock).filter(SiteBlock.block_key == "canva_saved_deck").first()
    if block and block.block_value:
        import json
        try:
            return json.loads(block.block_value)
        except Exception:
            pass
    return {
        "deck_title": "新竹青年分團 2026 年度成果簡報",
        "slides": [
            {
                "id": "slide_1",
                "template": "evaluation",
                "theme": "amber",
                "title": "四大宗旨落實成效：以文化弘揚佛法",
                "subtitle": "《新竹佛光青年X竹光三參》Podcast 與文藝沙龍",
                "bullets": [
                    "本年度累積製播 12 集節目，突破 12,000 次青年收聽人次",
                    "將《星雲說喻》經典寓言生活化，引導竹科青年轉化職場壓力",
                    "舉辦青年墨寶互動文創策展，總吸引 650 位大專青年與民眾參與"
                ],
                "metrics": [
                    {"num": "12,000+", "label": "線上收聽累積人次"},
                    {"num": "98%", "label": "活動滿意度問卷回饋"}
                ],
                "badge": "文化弘法・三好實踐",
                "footer": "國際佛光會新竹青年分團 成果評鑑委員會 報告"
            },
            {
                "id": "slide_2",
                "template": "promo",
                "theme": "emerald",
                "title": "2026 VEGRUN 全民蔬食路跑・青年活力打卡",
                "subtitle": "你跑步・我種樹・用雙腳守護綠色地球",
                "bullets": [
                    "號召大專青年與社區大眾，每日健走跑步記錄里程",
                    "結合十八尖山、青草湖與南寮漁港風景步道，推廣健康生活",
                    "透過數位後台管理系統，即時登錄並累積全團總里程"
                ],
                "metrics": [
                    {"num": "520 km", "label": "全團累積路跑里程"},
                    {"num": "86 位", "label": "青年志工熱情參與"}
                ],
                "badge": "VEGRUN・綠色地球",
                "footer": "身心靈環保・推廣蔬食愛地球"
            }
        ]
    }

@router.post("/save-deck")
def save_deck(deck: Dict[str, Any] = Body(...), db: Session = Depends(get_db), user: User = Depends(require_cadre)):
    import json
    json_str = json.dumps(deck, ensure_ascii=False)
    block = db.query(SiteBlock).filter(SiteBlock.block_key == "canva_saved_deck").first()
    if block:
        block.block_value = json_str
    else:
        block = SiteBlock(block_key="canva_saved_deck", block_title="Canva 編輯器草稿", block_value=json_str, category="slides")
    db.add(block)
    db.commit()
    return {"status": "success", "message": "簡報草稿已儲存"}

@router.post("/export-pptx")
def export_pptx(req: DeckExportRequest, current_user: User = Depends(require_cadre)):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    color_map = {
        "amber": {"bg": RGBColor(254, 252, 243), "primary": RGBColor(217, 119, 6), "secondary": RGBColor(180, 83, 9), "text": RGBColor(30, 41, 59), "card_bg": RGBColor(254, 243, 199), "accent": RGBColor(245, 158, 11)},
        "emerald": {"bg": RGBColor(240, 253, 244), "primary": RGBColor(5, 150, 105), "secondary": RGBColor(4, 120, 87), "text": RGBColor(30, 41, 59), "card_bg": RGBColor(209, 250, 229), "accent": RGBColor(16, 185, 129)},
        "blue": {"bg": RGBColor(248, 250, 252), "primary": RGBColor(37, 99, 235), "secondary": RGBColor(29, 78, 216), "text": RGBColor(30, 41, 59), "card_bg": RGBColor(224, 231, 255), "accent": RGBColor(59, 130, 246)},
        "earth": {"bg": RGBColor(255, 251, 245), "primary": RGBColor(161, 98, 7), "secondary": RGBColor(133, 77, 14), "text": RGBColor(30, 41, 59), "card_bg": RGBColor(254, 240, 199), "accent": RGBColor(202, 138, 4)}
    }

    for page in req.slides:
        slide = prs.slides.add_slide(blank_layout)
        theme = color_map.get(page.theme, color_map["amber"])

        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = theme["bg"]
        bg.line.fill.background()

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.18))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = theme["primary"]
        top_bar.line.fill.background()

        if page.badge:
            b_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(5.0), Inches(0.45))
            p = b_box.text_frame.paragraphs[0]
            p.text = f"★ {page.badge}"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = theme["primary"]

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.5), Inches(0.9))
        p = t_box.text_frame.paragraphs[0]
        p.text = page.title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = theme["text"]

        if page.subtitle:
            s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.5), Inches(0.6))
            p = s_box.text_frame.paragraphs[0]
            p.text = page.subtitle
            p.font.size = Pt(16)
            p.font.color.rgb = theme["secondary"]

        lw = Inches(7.5) if page.metrics else Inches(11.5)
        blist = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), lw, Inches(3.8))
        tf_b = blist.text_frame
        tf_b.word_wrap = True
        for idx, item in enumerate(page.bullets):
            p = tf_b.add_paragraph() if idx > 0 else tf_b.paragraphs[0]
            p.text = f"✔  {item}"
            p.font.size = Pt(15)
            p.font.color.rgb = theme["text"]
            p.space_after = Pt(16)

        if page.metrics:
            m_top = Inches(2.6)
            for m in page.metrics[:2]:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), m_top, Inches(3.7), Inches(1.6))
                card.fill.solid()
                card.fill.fore_color.rgb = theme["card_bg"]
                card.line.color.rgb = theme["accent"]
                p_num = card.text_frame.paragraphs[0]
                p_num.text = m.get("num", "0")
                p_num.alignment = PP_ALIGN.CENTER
                p_num.font.size = Pt(30)
                p_num.font.bold = True
                p_num.font.color.rgb = theme["primary"]
                p_lbl = card.text_frame.add_paragraph()
                p_lbl.text = m.get("label", "")
                p_lbl.alignment = PP_ALIGN.CENTER
                p_lbl.font.size = Pt(13)
                p_lbl.font.color.rgb = theme["text"]
                m_top += Inches(1.9)

        if page.footer:
            f_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.4))
            p = f_box.text_frame.paragraphs[0]
            p.text = page.footer
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(148, 163, 184)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"presentation_{timestamp}_{uuid.uuid4().hex[:4]}.pptx"
    out_path = UPLOAD_DIR / safe_name
    prs.save(str(out_path))

    return {
        "status": "success",
        "file_url": f"/uploads/{safe_name}",
        "filename": safe_name,
        "title": req.deck_title,
        "slides_count": len(req.slides)
    }
