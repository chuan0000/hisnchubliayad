import os
import shutil
import uuid
import json
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, Body, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel
try:
    from sqlalchemy.orm import Session
except ImportError:
    Session = object

from app.config import UPLOAD_DIR, ALLOWED_EXTENSIONS, MAX_FILE_SIZE_MB
from app.database import get_db
from app.models import User, SiteBlock
from app.auth import require_cadre

router = APIRouter(prefix="/api/upload", tags=["檔案上傳與本機存取 (File Upload)"])
slides_router = APIRouter(prefix="/api/slides", tags=["Canva 簡報與視覺編輯器 (Canva-like PPT Studio)"])

# ----------------- 【功能 7】自動浮水印壓印函式 (Pillow) -----------------
def add_watermark_to_image(image_path: Path, watermark_text: str = "國際佛光會新竹青年分團"):
    """使用 Pillow 為活動佐證照片壓印優雅半透明浮水印，保護青年肖像與活動資產"""
    try:
        from PIL import Image, ImageDraw, ImageFont

        font_path = None
        for p in [
            '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
            '/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc',
            '/System/Library/Fonts/PingFang.ttc',
            'C:\\Windows\\Fonts\\msjh.ttc'
        ]:
            if os.path.exists(p):
                font_path = p
                break

        with Image.open(image_path) as base_img:
            base_rgba = base_img.convert("RGBA")
            w, h = base_rgba.size

            font_size = max(16, int(h * 0.035))
            font = ImageFont.truetype(font_path, font_size) if font_path else ImageFont.load_default()

            full_text = f"☸ {watermark_text} ｜ 三好四給實踐存檔"
            overlay = Image.new("RGBA", (w, h), (255, 255, 255, 0))
            draw = ImageDraw.Draw(overlay)

            bbox = draw.textbbox((0, 0), full_text, font=font)
            text_w = bbox[2] - bbox[0]
            text_h = bbox[3] - bbox[1]

            margin_x = int(w * 0.03)
            margin_y = int(h * 0.03)
            x = w - text_w - margin_x
            y = h - text_h - margin_y

            padding_x = 12
            padding_y = 6
            # 深藍半透明藥丸底框
            draw.rounded_rectangle(
                [x - padding_x, y - padding_y, x + text_w + padding_x, y + text_h + padding_y],
                radius=8,
                fill=(15, 23, 42, 170)
            )
            # 佛光金文字
            draw.text((x, y), full_text, font=font, fill=(251, 191, 36, 240))

            res = Image.alpha_composite(base_rgba, overlay)
            if image_path.suffix.lower() in [".jpg", ".jpeg"]:
                res.convert("RGB").save(image_path, "JPEG", quality=92)
            else:
                res.save(image_path)
    except Exception as e:
        print(f"Watermark notice: {e}")

# ----------------- 檔案上傳與下載 API -----------------
@router.post("")
async def upload_file(
    file: UploadFile = File(...),
    watermark: bool = Query(True, description="【功能 7】是否自動為活動照片壓印分團浮水印"),
    current_user: User = Depends(require_cadre)
):
    """【活動幹部/管理員】上傳佐證照片或 PDF 公文至本機 ./uploads"""
    file_ext = Path(file.filename).suffix.lower()
    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"不支援的檔案格式: {file_ext}。僅支援圖片 (.jpg, .png, .webp) 與公文 (.pdf, .docx)"
        )

    time_prefix = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_id = uuid.uuid4().hex[:6]
    clean_name = Path(file.filename).name.replace(" ", "_")
    safe_filename = f"{time_prefix}_{unique_id}_{clean_name}"
    target_path = UPLOAD_DIR / safe_filename

    try:
        with target_path.open("wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"儲存檔案至伺服器失敗: {str(e)}")

    file_size = target_path.stat().st_size
    if file_size > MAX_FILE_SIZE_MB * 1024 * 1024:
        target_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"檔案過大，上限為 {MAX_FILE_SIZE_MB}MB")

    file_type = "pdf" if file_ext == ".pdf" else "image"

    # 【功能 49】提取相片原始 EXIF 拍攝時間與相機型號 (在壓印浮水印之前解析)
    exif_date = None
    camera_model = None
    if file_type == "image":
        try:
            from PIL import Image, ExifTags
            with Image.open(target_path) as img:
                exif = img.getexif()
                if exif:
                    for tag_id, val in exif.items():
                        tag_name = ExifTags.TAGS.get(tag_id, tag_id)
                        if tag_name in ("DateTimeOriginal", "DateTime"):
                            exif_date = str(val)[:10].replace(":", "-")
                        elif tag_name == "Model":
                            camera_model = str(val)
        except Exception:
            pass

    # 【功能 7】若為圖片且開啟浮水印，自動壓印
    is_watermarked = False
    if file_type == "image" and watermark:
        add_watermark_to_image(target_path)
        is_watermarked = True

    return {
        "exif_date": exif_date,
        "camera_model": camera_model,
        "status": "success",
        "file_url": f"/uploads/{safe_filename}",
        "saved_filename": safe_filename,
        "original_filename": file.filename,
        "file_type": file_type,
        "size_bytes": file_size,
        "watermarked": is_watermarked,
        "uploaded_by": current_user.full_name
    }

@router.get("/download/{filename}")
def download_file(filename: str):
    """下載指定檔案至個人電腦"""
    file_path = UPLOAD_DIR / filename
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(status_code=404, detail="檔案不存在")
    return FileResponse(path=file_path, filename=filename, media_type="application/octet-stream")

# ==============================================================================
# 【功能 49】照片 EXIF 拍攝時間與中繼資料解析 API
# ==============================================================================

# ==============================================================================
# 【功能 77】佛光三好祈福卡社群分享圖片生成器 API
# ==============================================================================
@router.get("/social-share-card")
def generate_social_share_card(
    quote: str = "做好事・說好話・存好心",
    author: str = "星雲大師法語",
    youth_name: str = "新竹佛光青年",
    theme: str = "golden"
):
    """【功能 77】產製名片/限時動態級三好祈福社群分享圖卡 (向量 SVG)"""
    themes = {
        "golden": {
            "bg_gradient": "linear-gradient(135deg, #78350F 0%, #1E293B 100%)",
            "primary": "#F59E0B",
            "secondary": "#FDE68A",
            "border": "#D97706",
            "badge_bg": "rgba(245, 158, 11, 0.2)"
        },
        "emerald": {
            "bg_gradient": "linear-gradient(135deg, #064E3B 0%, #0F172A 100%)",
            "primary": "#10B981",
            "secondary": "#A7F3D0",
            "border": "#059669",
            "badge_bg": "rgba(16, 185, 129, 0.2)"
        },
        "indigo": {
            "bg_gradient": "linear-gradient(135deg, #312E81 0%, #0F172A 100%)",
            "primary": "#6366F1",
            "secondary": "#C7D2FE",
            "border": "#4F46E5",
            "badge_bg": "rgba(99, 102, 241, 0.2)"
        }
    }
    th = themes.get(theme, themes["golden"])

    svg_markup = f"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 800" width="100%" height="100%" style="font-family: sans-serif; background: {th['bg_gradient']};">
      <rect x="24" y="24" width="552" height="752" rx="24" fill="none" stroke="{th['border']}" stroke-width="2" opacity="0.6" />
      <rect x="36" y="36" width="528" height="728" rx="18" fill="none" stroke="{th['border']}" stroke-width="1" stroke-dasharray="4,4" opacity="0.3" />

      <circle cx="300" cy="110" r="32" fill="{th['primary']}" opacity="0.15" />
      <circle cx="300" cy="110" r="24" fill="{th['primary']}" opacity="0.3" />
      <text x="300" y="117" font-size="20" fill="{th['primary']}" text-anchor="middle" font-weight="bold">☸</text>

      <text x="300" y="170" font-size="12" fill="{th['secondary']}" text-anchor="middle" font-weight="bold" letter-spacing="3">國際佛光會中華青年總會・新竹青年分團</text>
      <text x="300" y="195" font-size="10" fill="#94A3B8" text-anchor="middle" letter-spacing="2">THREE ACTS OF GOODNESS DAILY INSPIRATION</text>

      <line x1="200" y1="220" x2="400" y2="220" stroke="{th['primary']}" stroke-width="1.5" opacity="0.6" />

      <rect x="230" y="260" width="140" height="30" rx="15" fill="{th['badge_bg']}" stroke="{th['primary']}" stroke-width="1" />
      <text x="300" y="280" font-size="11" font-weight="bold" fill="{th['primary']}" text-anchor="middle">★ 三好四給行願卡 ★</text>

      <foreignObject x="60" y="320" width="480" height="200">
        <div xmlns="http://www.w3.org/1999/xhtml" style="font-size: 22px; line-height: 1.8; color: #F8FAFC; text-align: center; font-weight: bold; font-family: serif; text-shadow: 0 2px 4px rgba(0,0,0,0.5);">
          「{quote}」
        </div>
      </foreignObject>

      <text x="300" y="550" font-size="14" fill="{th['secondary']}" text-anchor="middle" font-weight="600">—— {author}</text>
      <text x="300" y="600" font-size="28" fill="{th['primary']}" text-anchor="middle" opacity="0.8">🪷</text>

      <rect x="60" y="650" width="480" height="60" rx="16" fill="rgba(0,0,0,0.3)" border="1px solid rgba(255,255,255,0.1)" />
      <text x="80" y="685" font-size="12" fill="#E2E8F0" font-weight="bold">發願實踐人：{youth_name}</text>
      <text x="520" y="685" font-size="11" fill="{th['secondary']}" text-anchor="end">佛光山新竹法寶寺 敬祝</text>
    </svg>"""

    return {
        "status": "success",
        "quote": quote,
        "author": author,
        "youth_name": youth_name,
        "theme": theme,
        "svg_markup": svg_markup
    }

@router.get("/extract-exif")
def extract_photo_exif(filename: str):
    """【功能 49】解析已上傳照片之 EXIF 拍攝時間，用於自動校正活動日期"""
    file_path = UPLOAD_DIR / Path(filename).name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="找不到指定照片檔案")

    exif_info = {"has_exif": False, "exif_date": None, "camera_model": None}
    try:
        from PIL import Image, ExifTags
        with Image.open(file_path) as img:
            exif = img.getexif()
            if exif:
                for tag_id, val in exif.items():
                    tag_name = ExifTags.TAGS.get(tag_id, tag_id)
                    if tag_name in ("DateTimeOriginal", "DateTime"):
                        exif_info["exif_date"] = str(val)[:10].replace(":", "-")
                        exif_info["has_exif"] = True
                    elif tag_name == "Model":
                        exif_info["camera_model"] = str(val)
    except Exception as e:
        exif_info["error"] = str(e)

    return exif_info


# ==============================================================================
# 【功能 52】活動相片自動美膚與佛光金暖色濾鏡預處理 API
# ==============================================================================
def apply_buddhist_light_filter(input_path: Path, output_path: Path, filter_type: str = "golden_warmth", intensity: float = 1.0):
    from PIL import Image, ImageEnhance, ImageFilter
    with Image.open(input_path) as img:
        orig = img.convert("RGB")
        # 1. 微妙肌膚柔光層
        blurred = orig.filter(ImageFilter.GaussianBlur(radius=2.0 * intensity))
        smoothed = Image.blend(orig, blurred, 0.15 * min(intensity, 1.5))
        # 2. 佛光暖金漸層遮罩 (255, 230, 160)
        gold_overlay = Image.new("RGB", orig.size, (255, 230, 160))
        golden_img = Image.blend(smoothed, gold_overlay, 0.12 * min(intensity, 1.5))
        # 3. 亮度微增通透 (+6%)
        bright = ImageEnhance.Brightness(golden_img).enhance(1.0 + 0.06 * intensity)
        # 4. 色彩飽和度提升，氣色紅潤 (+8%)
        color = ImageEnhance.Color(bright).enhance(1.0 + 0.08 * intensity)
        # 5. 清晰微對比 (+4%)
        contrast = ImageEnhance.Contrast(color).enhance(1.0 + 0.04 * intensity)
        contrast.save(str(output_path), "JPEG", quality=92)
    return output_path

class PhotoFilterRequest(BaseModel):
    filename: str
    filter_type: Optional[str] = "golden_warmth"
    intensity: Optional[float] = 1.0

def _process_filter(filename: str, filter_type: str, intensity: float):
    clean_name = Path(filename).name
    file_path = UPLOAD_DIR / clean_name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="找不到指定照片檔案")
    
    stem = file_path.stem
    out_name = f"golden_{stem}.jpg"
    out_path = UPLOAD_DIR / out_name
    try:
        apply_buddhist_light_filter(file_path, out_path, filter_type=filter_type, intensity=intensity)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"濾鏡處理失敗: {str(e)}")
    
    return {
        "status": "success",
        "original_filename": clean_name,
        "filtered_filename": out_name,
        "url": f"/uploads/{out_name}",
        "filter_type": filter_type,
        "filter_name": "佛光金暖光柔膚濾鏡 (Golden Warmth & Skin Glow)",
        "enhancements": {
            "color_temperature": "暖金調色 +12% (#FFE6A0 暖金色調)",
            "skin_smoothing": "高斯柔光高光層合成 (15% 輕量柔膚)",
            "brightness": "+6% 提亮通透法喜充滿",
            "saturation": "+8% 氣色紅潤精神煥發"
        }
    }

@router.post("/apply-filter")
def apply_photo_filter_post(req: PhotoFilterRequest):
    """【功能 52】對活動照片套用佛光金暖色柔膚濾鏡 (POST)"""
    return _process_filter(req.filename, req.filter_type or "golden_warmth", req.intensity or 1.0)

@router.get("/apply-filter")
def apply_photo_filter_get(filename: str, filter_type: str = "golden_warmth", intensity: float = 1.0):
    """【功能 52】對活動照片套用佛光金暖色柔膚濾鏡 (GET)"""
    return _process_filter(filename, filter_type, intensity)

# ----------------- Canva 簡報編輯器與 PPTX 匯出 API -----------------
class SlidePage(BaseModel):
    id: Optional[str] = None
    template: str = "evaluation"
    theme: str = "amber"
    title: str = "四大宗旨活動成果評鑑"
    subtitle: str = "國際佛光會新竹青年分團"
    bullets: List[str] = []
    metrics: List[Dict[str, str]] = []
    badge: str = "三好四給・生活實踐"
    footer: str = "做好事・說好話・存好心"
    bg_image: Optional[str] = ""

class DeckExportRequest(BaseModel):
    deck_title: str = "新竹佛光青年分團簡報"
    slides: List[SlidePage]

@slides_router.get("/get-deck")
def get_saved_deck(db: Session = Depends(get_db)):
    block = db.query(SiteBlock).filter(SiteBlock.block_key == "canva_saved_deck").first()
    if block and block.block_value:
        try:
            return json.loads(block.block_value)
        except Exception:
            pass
    
    return {
        "deck_title": "新竹青年分團 2026 年度成果簡報",
        "slides": [
            {
                "id": "slide_1",
                "template": "evaluation",
                "theme": "amber",
                "title": "四大宗旨落實成效：以文化弘揚佛法",
                "subtitle": "《新竹佛光青年X竹光三參》Podcast 與文藝沙龍",
                "bullets": [
                    "本年度累積製播 12 集節目，突破 12,000 次青年收聽人次",
                    "將《星雲說喻》經典寓言生活化，引導竹科青年轉化職場壓力",
                    "舉辦青年墨寶互動文創策展，總吸引 650 位大專青年與民眾參與"
                ],
                "metrics": [
                    {"num": "12,000+", "label": "線上收聽累積人次"},
                    {"num": "98%", "label": "活動滿意度問卷回饋"}
                ],
                "badge": "文化弘法・三好實踐",
                "footer": "國際佛光會新竹青年分團 成果評鑑委員會 報告"
            },
            {
                "id": "slide_2",
                "template": "promo",
                "theme": "emerald",
                "title": "2026 VEGRUN 全民蔬食路跑・青年活力打卡",
                "subtitle": "你跑步・我種樹・用雙腳守護綠色地球",
                "bullets": [
                    "號召大專青年與社區大眾，每日健走跑步記錄里程",
                    "結合十八尖山、青草湖與南寮漁港風景步道，推廣健康生活",
                    "透過數位後台管理系統，即時登錄並累積全團總里程"
                ],
                "metrics": [
                    {"num": "520 km", "label": "全團累積路跑里程"},
                    {"num": "86 位", "label": "青年志工熱情參與"}
                ],
                "badge": "VEGRUN・綠色地球",
                "footer": "身心靈環保・推廣蔬食愛地球"
            }
        ]
    }

@slides_router.post("/save-deck")
def save_deck(
    deck: Dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_cadre)
):
    json_str = json.dumps(deck, ensure_ascii=False)
    block = db.query(SiteBlock).filter(SiteBlock.block_key == "canva_saved_deck").first()
    if block:
        block.block_value = json_str
    else:
        block = SiteBlock(
            block_key="canva_saved_deck",
            block_title="Canva 編輯器投影片草稿",
            block_value=json_str,
            category="slides",
            value_type="markdown"
        )
    db.add(block)
    db.commit()
    return {"status": "success", "message": "簡報草稿已儲存"}

@slides_router.post("/export-pptx")
def export_pptx(
    req: DeckExportRequest,
    current_user: User = Depends(require_cadre)
):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    color_map = {
        "amber": {
            "bg": RGBColor(254, 252, 243),
            "primary": RGBColor(217, 119, 6),
            "secondary": RGBColor(180, 83, 9),
            "text": RGBColor(30, 41, 59),
            "card_bg": RGBColor(254, 243, 199),
            "accent": RGBColor(245, 158, 11)
        },
        "emerald": {
            "bg": RGBColor(240, 253, 244),
            "primary": RGBColor(5, 150, 105),
            "secondary": RGBColor(4, 120, 87),
            "text": RGBColor(30, 41, 59),
            "card_bg": RGBColor(209, 250, 229),
            "accent": RGBColor(16, 185, 129)
        },
        "blue": {
            "bg": RGBColor(248, 250, 252),
            "primary": RGBColor(37, 99, 235),
            "secondary": RGBColor(29, 78, 216),
            "text": RGBColor(30, 41, 59),
            "card_bg": RGBColor(224, 231, 255),
            "accent": RGBColor(59, 130, 246)
        },
        "earth": {
            "bg": RGBColor(255, 251, 245),
            "primary": RGBColor(161, 98, 7),
            "secondary": RGBColor(133, 77, 14),
            "text": RGBColor(30, 41, 59),
            "card_bg": RGBColor(254, 240, 199),
            "accent": RGBColor(202, 138, 4)
        }
    }

    for page in req.slides:
        slide = prs.slides.add_slide(blank_slide_layout)
        theme = color_map.get(page.theme, color_map["amber"])

        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = theme["bg"]
        bg_shape.line.fill.background()

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.18))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = theme["primary"]
        top_bar.line.fill.background()

        if page.badge:
            badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(5.0), Inches(0.45))
            tf_badge = badge_box.text_frame
            tf_badge.word_wrap = True
            p_badge = tf_badge.paragraphs[0]
            p_badge.text = f"★ {page.badge}"
            p_badge.font.size = Pt(13)
            p_badge.font.bold = True
            p_badge.font.color.rgb = theme["primary"]

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.5), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = page.title
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = theme["text"]

        if page.subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.5), Inches(0.6))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = page.subtitle
            p_sub.font.size = Pt(16)
            p_sub.font.color.rgb = theme["secondary"]

        left_w = Inches(7.5) if page.metrics else Inches(11.5)
        bullets_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), left_w, Inches(3.8))
        tf_bullets = bullets_box.text_frame
        tf_bullets.word_wrap = True
        
        for idx, bullet_text in enumerate(page.bullets):
            p = tf_bullets.add_paragraph() if idx > 0 else tf_bullets.paragraphs[0]
            p.text = f"✔  {bullet_text}"
            p.font.size = Pt(15)
            p.font.color.rgb = theme["text"]
            p.space_after = Pt(16)

        if page.metrics:
            m_top = Inches(2.6)
            for m in page.metrics[:2]:
                card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), m_top, Inches(3.7), Inches(1.6))
                card_shape.fill.solid()
                card_shape.fill.fore_color.rgb = theme["card_bg"]
                card_shape.line.color.rgb = theme["accent"]
                card_shape.line.width = Pt(1.5)

                tf_card = card_shape.text_frame
                tf_card.word_wrap = True
                
                p_num = tf_card.paragraphs[0]
                p_num.text = m.get("num", "0")
                p_num.alignment = PP_ALIGN.CENTER
                p_num.font.size = Pt(30)
                p_num.font.bold = True
                p_num.font.color.rgb = theme["primary"]

                p_lbl = tf_card.add_paragraph()
                p_lbl.text = m.get("label", "")
                p_lbl.alignment = PP_ALIGN.CENTER
                p_lbl.font.size = Pt(13)
                p_lbl.font.color.rgb = theme["text"]

                m_top += Inches(1.9)

        if page.footer:
            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.4))
            tf_foot = footer_box.text_frame
            p_foot = tf_foot.paragraphs[0]
            p_foot.text = page.footer
            p_foot.font.size = Pt(11)
            p_foot.font.color.rgb = RGBColor(148, 163, 184)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"presentation_{timestamp}_{uuid.uuid4().hex[:4]}.pptx"
    out_path = UPLOAD_DIR / safe_name
    prs.save(str(out_path))

    return {
        "status": "success",
        "file_url": f"/uploads/{safe_name}",
        "filename": safe_name,
        "title": req.deck_title,
        "slides_count": len(req.slides)
    }
